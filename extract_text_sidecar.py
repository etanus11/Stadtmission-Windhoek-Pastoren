import os
from pptx import Presentation

prs = Presentation('StamiPastoren.pptx')

for i, slide in enumerate(prs.slides, start=1):
    text_content = []
    
    # Check slide text frames
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text.strip():
            text_content.append(shape.text.strip())
            
    # Check speaker notes (fallback/addition)
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            text_content.append(f"Notes: {notes}")

    # Write sidecar file if text exists
    if text_content:
        filename = f"slide-{i:02d}.txt"
        with open(filename, "w", encoding="utf-8") as f:
            f.write("\n".join(text_content))
        print(f"Created {filename}")